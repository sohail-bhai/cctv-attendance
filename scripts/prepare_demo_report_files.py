"""
Prepare demo files for the React Report Analyzer.

This script copies the latest generated attendance report CSVs into:
    frontend/public/demo_reports/

The React Report Analyzer can then load those files with one button click.

Works with both:
1) normal CSV outputs from mark_attendance_checkpoints.py
2) an XLSX-only report, by exporting Excel sheets back to CSV when openpyxl is available
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "frontend" / "public" / "demo_reports"

KINDS = {
    "attendance": "attendance",
    "slot_summary": "slot_summary",
    "detection_log": "detection_log",
    "checkpoint_summary": "checkpoint_summary",
    "student_checkpoint_evidence": "student_checkpoint_evidence",
    "camera_summary": "camera_summary",
    "needs_review": "needs_review",
}

SHEET_TO_KIND = {
    "Final Attendance": "attendance",
    "Attendance": "attendance",
    "Slot Summary": "slot_summary",
    "Detection Log": "detection_log",
    "Checkpoint Summary": "checkpoint_summary",
    "Student Checkpoints": "student_checkpoint_evidence",
    "Student Checkpoint Evidence": "student_checkpoint_evidence",
    "Camera Summary": "camera_summary",
    "Needs Review": "needs_review",
}


@dataclass
class Candidate:
    path: Path
    mtime: float


def clean_name(value: str) -> str:
    value = value.strip().replace(" ", "_")
    value = re.sub(r"[^A-Za-z0-9_\-]+", "", value)
    return value or "demo"


def iter_files(source: Path) -> Iterable[Path]:
    if not source.exists():
        return []
    if source.is_file():
        return [source]
    return [p for p in source.rglob("*") if p.is_file()]


def infer_kind(path: Path) -> Optional[str]:
    lower = path.name.lower()
    # Longest names first so student_checkpoint_evidence does not become checkpoint_summary.
    for kind in sorted(KINDS.keys(), key=len, reverse=True):
        if lower.startswith(kind.lower()) or f"{kind.lower()}_" in lower:
            return kind
    return None


def find_latest_csvs(source: Path, slot_id: str | None) -> dict[str, Path]:
    candidates: dict[str, list[Candidate]] = {kind: [] for kind in KINDS}
    slot_text = (slot_id or "").lower()

    for path in iter_files(source):
        if path.suffix.lower() != ".csv":
            continue
        if slot_text and slot_text not in path.name.lower():
            # still allow nested folder names containing slot_id
            if slot_text not in str(path.parent).lower():
                continue
        kind = infer_kind(path)
        if kind:
            candidates[kind].append(Candidate(path=path, mtime=path.stat().st_mtime))

    latest: dict[str, Path] = {}
    for kind, items in candidates.items():
        if items:
            latest[kind] = max(items, key=lambda item: item.mtime).path
    return latest


def find_latest_xlsx(source: Path, slot_id: str | None) -> Optional[Path]:
    slot_text = (slot_id or "").lower()
    candidates = []
    for path in iter_files(source):
        if path.suffix.lower() != ".xlsx":
            continue
        if slot_text and slot_text not in path.name.lower() and slot_text not in str(path.parent).lower():
            continue
        candidates.append(path)
    if not candidates:
        return None
    return max(candidates, key=lambda p: p.stat().st_mtime)


def copy_csvs(files: dict[str, Path], output_dir: Path, prefix: str) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    public_paths: dict[str, str] = {}
    for kind, src in files.items():
        dest = output_dir / f"{prefix}_{kind}.csv"
        shutil.copy2(src, dest)
        public_paths[kind] = f"/demo_reports/{dest.name}"
    return public_paths


def export_xlsx_to_csv(xlsx_path: Path, output_dir: Path, prefix: str) -> dict[str, str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit(
            "openpyxl is required to export XLSX demo files. Install it with: pip install openpyxl"
        ) from exc

    output_dir.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    public_paths: dict[str, str] = {}

    for sheet_name in wb.sheetnames:
        kind = SHEET_TO_KIND.get(sheet_name)
        if not kind:
            continue
        ws = wb[sheet_name]
        dest = output_dir / f"{prefix}_{kind}.csv"
        with dest.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(["" if cell is None else cell for cell in row])
        public_paths[kind] = f"/demo_reports/{dest.name}"
    return public_paths


def write_manifest(output_dir: Path, prefix: str, files: dict[str, str], source: Path) -> Path:
    manifest = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "title": "MON_P1 Demo Report",
        "description": "Demo files generated from the latest existing attendance output.",
        "source": str(source),
        "files": files,
        "notes": [
            "Use this only for Report Analyzer demo/testing.",
            "It does not rerun attendance processing.",
            "Regenerate after every new report if you want the demo button to show the latest data.",
        ],
    }
    manifest_path = output_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepare demo files for React Report Analyzer buttons.")
    parser.add_argument("--source", default="attendance_output", help="Folder or file to scan. Default: attendance_output")
    parser.add_argument("--slot-id", default="MON_P1", help="Slot id to prefer/filter, e.g. MON_P1")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Output folder. Default: frontend/public/demo_reports")
    parser.add_argument("--prefix", default="mon_p1_demo", help="Prefix for generated demo CSV filenames")
    parser.add_argument("--from-xlsx", default="", help="Optional exact XLSX file path to export to demo CSVs")
    args = parser.parse_args()

    source = (PROJECT_ROOT / args.source).resolve() if not Path(args.source).is_absolute() else Path(args.source)
    output_dir = (PROJECT_ROOT / args.output_dir).resolve() if not Path(args.output_dir).is_absolute() else Path(args.output_dir)
    prefix = clean_name(args.prefix)

    if args.from_xlsx:
        xlsx = (PROJECT_ROOT / args.from_xlsx).resolve() if not Path(args.from_xlsx).is_absolute() else Path(args.from_xlsx)
        if not xlsx.exists():
            raise SystemExit(f"XLSX file not found: {xlsx}")
        files = export_xlsx_to_csv(xlsx, output_dir, prefix)
        source_used = xlsx
    else:
        csvs = find_latest_csvs(source, args.slot_id)
        if csvs:
            files = copy_csvs(csvs, output_dir, prefix)
            source_used = source
        else:
            xlsx = find_latest_xlsx(source, args.slot_id)
            if not xlsx:
                raise SystemExit(
                    f"No CSV or XLSX report files found in {source}. Run attendance once first, or pass --from-xlsx path."
                )
            files = export_xlsx_to_csv(xlsx, output_dir, prefix)
            source_used = xlsx

    if "attendance" not in files:
        raise SystemExit("Could not prepare attendance CSV. Need at least an attendance report.")

    manifest_path = write_manifest(output_dir, prefix, files, source_used)

    print("Demo report files prepared successfully.")
    print(f"Output folder: {output_dir}")
    print(f"Manifest: {manifest_path}")
    print("Files:")
    for kind, public_path in files.items():
        print(f"  {kind}: {public_path}")


if __name__ == "__main__":
    main()
